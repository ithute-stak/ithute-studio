from __future__ import annotations

import base64
import io
import re
from copy import deepcopy
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdf_canvas

from app.document_renderer import render_docx as _render_docx
from app.document_renderer import render_pdf as _render_pdf
from app.html_renderer import render_html as _render_html

Json = dict[str, Any]


def _get_path(value: Any, *parts: str) -> Any:
    current = value
    for part in parts:
        if not isinstance(current, dict):
            return None
        current = current.get(part)
    return current


def _brand_name(document: Json) -> str:
    variables = document.get("variables") or {}
    for path in (
        ("company", "name"),
        ("organization", "legalName"),
        ("organization", "name"),
    ):
        value = _get_path(variables, *path)
        if value:
            return str(value)
    header = document.get("header") or {}
    for node in header.get("content") or []:
        text = "".join(
            str(child.get("text") or "")
            for child in node.get("content") or []
            if child.get("type") == "text"
        ).strip()
        if text:
            return text
    return "Ithute Studio"


def _decode_data_uri(value: str) -> bytes | None:
    match = re.match(r"^data:[^;,]+;base64,(.*)$", value, re.S)
    if not match:
        return None
    try:
        return base64.b64decode(match.group(1))
    except Exception:
        return None


def _logo_bytes(document: Json) -> bytes | None:
    variables = document.get("variables") or {}
    candidates = [
        _get_path(variables, "company", "logoUrl"),
        _get_path(variables, "company", "logo"),
        _get_path(variables, "company", "branding", "logoUrl"),
        _get_path(variables, "organization", "branding", "logoUrl"),
        (document.get("design") or {}).get("logoUrl"),
    ]
    for candidate in candidates:
        if not candidate:
            continue
        value = str(candidate)
        if "{{" in value:
            continue
        decoded = _decode_data_uri(value)
        if decoded:
            return decoded
    return None


def _safe_hex(value: object, fallback: str = "#0B5EA8") -> str:
    raw = str(value or fallback).strip()
    if re.fullmatch(r"#[0-9A-Fa-f]{6}", raw):
        return raw
    return fallback


def _initials(name: str) -> str:
    words = [part for part in re.split(r"\s+", name.strip()) if part]
    if not words:
        return "IS"
    if len(words) == 1:
        return words[0][:2].upper()
    return f"{words[0][0]}{words[-1][0]}".upper()


def _monogram_png(name: str, primary: str) -> bytes:
    size = 220
    image = Image.new("RGBA", (size, size), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    color = _safe_hex(primary)
    draw.rounded_rectangle((5, 5, size - 5, size - 5), radius=52, fill=color)
    draw.rounded_rectangle(
        (17, 17, size - 17, size - 17),
        radius=42,
        outline=(255, 255, 255, 150),
        width=5,
    )
    text = _initials(name)
    font = ImageFont.load_default(size=76)
    box = draw.textbbox((0, 0), text, font=font)
    width = box[2] - box[0]
    height = box[3] - box[1]
    draw.text(
        ((size - width) / 2, (size - height) / 2 - 6),
        text,
        fill="white",
        font=font,
    )
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _brand_image(document: Json) -> bytes:
    logo = _logo_bytes(document)
    if logo:
        try:
            with Image.open(io.BytesIO(logo)) as source:
                source.thumbnail((560, 220), Image.Resampling.LANCZOS)
                rendered = source.convert("RGBA")
                output = io.BytesIO()
                rendered.save(output, format="PNG", optimize=True)
                return output.getvalue()
        except Exception:
            pass
    design = document.get("design") or {}
    return _monogram_png(_brand_name(document), str(design.get("primaryColor") or "#0B5EA8"))


def _presentation_document(document: Json) -> Json:
    prepared = deepcopy(document)
    design = prepared.setdefault("design", {})
    if isinstance(design, dict):
        design.setdefault("logoUrl", "{{company.logoUrl}}")
        design.setdefault("brandMode", "logo-with-monogram-fallback")
        design.setdefault("visualProfile", "premium-2026-09")
    return prepared


def render_pdf(document: Json) -> bytes:
    prepared = _presentation_document(document)
    base = _render_pdf(prepared)
    logo = _brand_image(prepared)
    reader = PdfReader(io.BytesIO(base))
    writer = PdfWriter()

    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_buffer = io.BytesIO()
        overlay = pdf_canvas.Canvas(overlay_buffer, pagesize=(width, height))
        style = str((prepared.get("design") or {}).get("headerStyle") or "minimal")
        x = 8.5 * mm if style == "left-accent" else 2.8 * mm
        y = height - 18.5 * mm
        box = 12.5 * mm
        overlay.setFillColor(colors.white)
        overlay.roundRect(x - 0.8 * mm, y - 0.8 * mm, box + 1.6 * mm, box + 1.6 * mm, 2.7 * mm, fill=1, stroke=0)
        overlay.drawImage(
            ImageReader(io.BytesIO(logo)),
            x,
            y,
            width=box,
            height=box,
            preserveAspectRatio=True,
            anchor="c",
            mask="auto",
        )
        overlay.save()
        overlay_buffer.seek(0)
        page.merge_page(PdfReader(overlay_buffer).pages[0], over=True)
        writer.add_page(page)

    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def render_docx(document: Json) -> bytes:
    prepared = _presentation_document(document)
    base = _render_docx(prepared)
    doc = DocxDocument(io.BytesIO(base))
    logo = _brand_image(prepared)

    for section in doc.sections:
        header = section.header
        brand = header.add_paragraph()
        brand.alignment = WD_ALIGN_PARAGRAPH.LEFT
        brand.paragraph_format.space_before = Pt(0)
        brand.paragraph_format.space_after = Pt(0)
        brand.paragraph_format.line_spacing = 1
        run = brand.add_run()
        run.add_picture(io.BytesIO(logo), width=Mm(13))
        header._element.remove(brand._p)
        header._element.insert(0, brand._p)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_html(document: Json) -> str:
    prepared = _presentation_document(document)
    rendered = _render_html(prepared)
    logo = _brand_image(prepared)
    encoded = base64.b64encode(logo).decode("ascii")
    brand = (
        '<div class="ithute-brand-lockup" style="position:fixed;top:8mm;left:3mm;z-index:10;'
        'width:12mm;height:12mm;padding:1mm;border-radius:3mm;background:#fff;'
        'box-shadow:0 2px 8px rgba(15,23,42,.12)">'
        f'<img src="data:image/png;base64,{encoded}" alt="Company logo" '
        'style="width:100%;height:100%;object-fit:contain"></div>'
    )
    return rendered.replace("<body>", f"<body>{brand}", 1)
