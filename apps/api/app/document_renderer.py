from __future__ import annotations

import base64
import io
import re
from html import escape
from typing import Any
from urllib.parse import unquote

import qrcode
from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as PdfImage,
    PageBreak as PdfPageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

Json = dict[str, Any]

DEFAULT_DESIGN: Json = {
    "primaryColor": "#1F2937",
    "secondaryColor": "#F3F4F6",
    "accentColor": "#6B7280",
    "textColor": "#111827",
    "mutedColor": "#6B7280",
    "borderColor": "#D1D5DB",
    "fontFamily": "Helvetica",
    "headerStyle": "minimal",
    "footerStyle": "minimal",
}
STANDARD_PDF_FONTS = {
    "Helvetica",
    "Helvetica-Bold",
    "Helvetica-Oblique",
    "Times-Roman",
    "Times-Bold",
    "Times-Italic",
    "Courier",
    "Courier-Bold",
    "Courier-Oblique",
}


def _attrs(node: Json) -> Json:
    return node.get("attrs") or {}


def _design(document: Json) -> Json:
    return {**DEFAULT_DESIGN, **(document.get("design") or {})}


def _plain_text(node: Json | None) -> str:
    if not node:
        return ""
    if node.get("type") == "text":
        return str(node.get("text", ""))
    if node.get("type") == "hardBreak":
        return "\n"
    return "".join(_plain_text(child) for child in node.get("content") or [])


def _region_lines(node: Json | None) -> list[str]:
    if not node:
        return []
    content = node.get("content") or []
    lines = [" ".join(_plain_text(item).split()) for item in content]
    return [line for line in lines if line]


def _decode_data_uri(uri: str) -> bytes | None:
    match = re.match(r"^data:([^;,]+)?(;base64)?,(.*)$", uri, re.S)
    if not match:
        return None
    payload = match.group(3)
    try:
        return base64.b64decode(payload) if match.group(2) else unquote(payload).encode()
    except Exception:
        return None


def _page_size(settings: Json):
    size = LETTER if settings.get("pageSize") == "letter" else A4
    return landscape(size) if settings.get("orientation") == "landscape" else size


def _margins(settings: Json):
    values = settings.get("marginsMm") or {}
    return tuple(float(values.get(key, 20)) * mm for key in ("top", "right", "bottom", "left"))


def _pdf_font(name: str, bold: bool = False) -> str:
    if name in STANDARD_PDF_FONTS:
        if bold and name == "Helvetica":
            return "Helvetica-Bold"
        if bold and name == "Times-Roman":
            return "Times-Bold"
        if bold and name == "Courier":
            return "Courier-Bold"
        return name
    lower = name.lower()
    base = "Times" if "times" in lower or "serif" in lower else "Helvetica"
    return f"{base}-Bold" if bold else ("Times-Roman" if base == "Times" else "Helvetica")


def _pdf_markup(node: Json) -> str:
    if node.get("type") == "text":
        text = escape(str(node.get("text", "")))
        for mark in node.get("marks") or []:
            kind = mark.get("type")
            if kind == "bold":
                text = f"<b>{text}</b>"
            elif kind == "italic":
                text = f"<i>{text}</i>"
            elif kind == "underline":
                text = f"<u>{text}</u>"
            elif kind == "link":
                href = escape(str((mark.get("attrs") or {}).get("href", "")))
                text = f'<link href="{href}">{text}</link>'
        return text
    if node.get("type") == "hardBreak":
        return "<br/>"
    return "".join(_pdf_markup(child) for child in node.get("content") or [])


def _pdf_alignment(node: Json):
    align = str(_attrs(node).get("textAlign") or "left")
    return {
        "left": TA_LEFT,
        "center": TA_CENTER,
        "right": TA_RIGHT,
        "justify": TA_JUSTIFY,
    }.get(align, TA_LEFT)


def _qr_image(value: str, width: float = 28 * mm):
    qr = qrcode.QRCode(box_size=6, border=2)
    qr.add_data(value or " ")
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    buf.seek(0)
    return PdfImage(buf, width=width, height=width)


def _pdf_story(nodes: list[Json], styles, design: Json) -> list[Any]:
    story: list[Any] = []
    primary = colors.HexColor(str(design["primaryColor"]))
    secondary = colors.HexColor(str(design["secondaryColor"]))
    text_color = colors.HexColor(str(design["textColor"]))
    border_color = colors.HexColor(str(design["borderColor"]))
    font = _pdf_font(str(design["fontFamily"]))

    for node in nodes:
        kind = node.get("type")
        attrs = _attrs(node)
        if kind in {"paragraph", "heading"}:
            if kind == "paragraph":
                base = styles["BodyText"]
                size = 9.5
                color = text_color
                font_name = font
                space_after = 4
            else:
                level = min(3, int(attrs.get("level", 1)))
                base = styles[f"Heading{level}"]
                size = {1: 20, 2: 13, 3: 11}[level]
                color = primary
                font_name = _pdf_font(str(design["fontFamily"]), bold=True)
                space_after = 7
            style = ParagraphStyle(
                f"ithute-{id(node)}",
                parent=base,
                alignment=_pdf_alignment(node),
                fontName=font_name,
                fontSize=size,
                leading=size * 1.25,
                textColor=color,
                leftIndent=float(attrs.get("indent") or 0) * 18,
                firstLineIndent=float(attrs.get("firstLineIndent") or 0) * mm,
                backColor=colors.HexColor(str(attrs["shading"])) if attrs.get("shading") else None,
                borderColor=border_color if attrs.get("border") else None,
                borderWidth=0.7 if attrs.get("border") else 0,
                borderPadding=5 if attrs.get("border") else 0,
                spaceAfter=space_after,
            )
            story.append(Paragraph(_pdf_markup(node) or "&nbsp;", style))
        elif kind == "pageBreak":
            story.append(PdfPageBreak())
        elif kind == "horizontalRule":
            story.append(
                Table(
                    [[""]],
                    colWidths=[170 * mm],
                    rowHeights=[1],
                    style=TableStyle([("LINEABOVE", (0, 0), (-1, -1), 0.7, primary)]),
                )
            )
        elif kind == "table":
            rows = []
            for row in node.get("content") or []:
                rows.append(
                    [
                        Paragraph(
                            _pdf_markup(cell) or "&nbsp;",
                            ParagraphStyle(
                                f"cell-{id(cell)}",
                                parent=styles["BodyText"],
                                fontName=font,
                                fontSize=8.5,
                                leading=10.5,
                                textColor=text_color,
                            ),
                        )
                        for cell in row.get("content") or []
                    ]
                )
            if rows:
                table = Table(rows, repeatRows=1, hAlign="LEFT")
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.45, border_color),
                            ("BACKGROUND", (0, 0), (-1, 0), primary),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, 0), _pdf_font(str(design["fontFamily"]), bold=True)),
                            ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, secondary]),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 5),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                            ("TOPPADDING", (0, 0), (-1, -1), 5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ]
                    )
                )
                story.append(table)
                story.append(Spacer(1, 7))
        elif kind == "image":
            src = str(attrs.get("src") or "")
            raw = _decode_data_uri(src)
            if raw:
                image = PdfImage(io.BytesIO(raw))
                image._restrictSize(170 * mm, 120 * mm)
                story.append(image)
                story.append(Spacer(1, 6))
        elif kind == "qrCode":
            value = str(attrs.get("value") or attrs.get("label") or "Ithute Document")
            story.append(_qr_image(value))
            story.append(Spacer(1, 6))
        elif kind in {"signature", "stamp", "logo", "letterhead", "chart"}:
            label = escape(str(attrs.get("label") or kind.title()))
            style = ParagraphStyle(
                f"placeholder-{id(node)}",
                parent=styles["BodyText"],
                textColor=primary,
                fontName=_pdf_font(str(design["fontFamily"]), bold=True),
            )
            story.append(Paragraph(f"<b>{label}</b>", style))
            story.append(Spacer(1, 12))
        elif kind == "textBox":
            inner = _pdf_story(node.get("content") or [], styles, design)
            if inner:
                story.append(
                    Table(
                        [[inner]],
                        style=TableStyle(
                            [
                                ("BOX", (0, 0), (-1, -1), 0.8, border_color),
                                ("BACKGROUND", (0, 0), (-1, -1), secondary),
                                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                ("TOPPADDING", (0, 0), (-1, -1), 8),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                            ]
                        ),
                    )
                )
        elif kind in {"bulletList", "orderedList"}:
            for index, item in enumerate(node.get("content") or [], 1):
                prefix = "•" if kind == "bulletList" else f"{index}."
                story.append(Paragraph(f"{prefix} {_pdf_markup(item)}", styles["BodyText"]))
        elif node.get("content"):
            story.extend(_pdf_story(node["content"], styles, design))
    return story


def _draw_pdf_header(canvas, width: float, height: float, left: float, right: float, lines: list[str], design: Json) -> None:
    primary = colors.HexColor(str(design["primaryColor"]))
    secondary = colors.HexColor(str(design["secondaryColor"]))
    accent = colors.HexColor(str(design["accentColor"]))
    muted = colors.HexColor(str(design["mutedColor"]))
    style = str(design.get("headerStyle") or "minimal")
    company = lines[0] if lines else ""
    label = lines[1] if len(lines) > 1 else str(design.get("documentLabel") or "")
    detail = lines[2] if len(lines) > 2 else ""

    if style == "band":
        canvas.setFillColor(primary)
        canvas.rect(0, height - 26 * mm, width, 26 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 12)
        canvas.drawString(left, height - 10 * mm, company[:85])
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 8)
        canvas.drawRightString(width - right, height - 10 * mm, label[:70])
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7.5)
        canvas.drawString(left, height - 17 * mm, detail[:100])
        canvas.setFillColor(accent)
        canvas.rect(0, height - 27 * mm, width, 1.3 * mm, fill=1, stroke=0)
    elif style == "split":
        canvas.setFillColor(primary)
        canvas.rect(0, height - 23 * mm, width * 0.68, 23 * mm, fill=1, stroke=0)
        canvas.setFillColor(secondary)
        canvas.rect(width * 0.68, height - 23 * mm, width * 0.32, 23 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 11)
        canvas.drawString(left, height - 10 * mm, company[:70])
        canvas.setFillColor(primary)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 8)
        canvas.drawRightString(width - right, height - 10 * mm, label[:55])
        canvas.setFillColor(accent)
        canvas.rect(width * 0.68, height - 24 * mm, width * 0.32, 1.3 * mm, fill=1, stroke=0)
    elif style == "left-accent":
        canvas.setFillColor(primary)
        canvas.rect(0, height - 28 * mm, 7 * mm, 28 * mm, fill=1, stroke=0)
        canvas.setFillColor(accent)
        canvas.rect(7 * mm, height - 3 * mm, width - 7 * mm, 1.3 * mm, fill=1, stroke=0)
        canvas.setFillColor(primary)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 12)
        canvas.drawString(left, height - 11 * mm, company[:85])
        canvas.setFillColor(accent)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 9)
        canvas.drawRightString(width - right, height - 11 * mm, label[:65])
        canvas.setFillColor(muted)
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7.5)
        canvas.drawString(left, height - 18 * mm, detail[:100])
    elif style == "ledger":
        canvas.setFillColor(primary)
        canvas.rect(0, height - 18 * mm, width, 18 * mm, fill=1, stroke=0)
        canvas.setFillColor(secondary)
        canvas.rect(0, height - 25 * mm, width, 7 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 11)
        canvas.drawString(left, height - 10.5 * mm, company[:85])
        canvas.drawRightString(width - right, height - 10.5 * mm, label[:65])
        canvas.setFillColor(primary)
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7.5)
        canvas.drawString(left, height - 22.5 * mm, detail[:110])
    else:
        canvas.setStrokeColor(primary)
        canvas.setLineWidth(1.2)
        canvas.line(left, height - 7 * mm, width - right, height - 7 * mm)
        canvas.setFillColor(primary)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 10.5)
        canvas.drawString(left, height - 13 * mm, company[:85])
        canvas.setFillColor(muted)
        canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 8)
        canvas.drawRightString(width - right, height - 13 * mm, label[:65])
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7.2)
        canvas.drawString(left, height - 19 * mm, detail[:100])


def _draw_pdf_footer(canvas, width: float, left: float, right: float, lines: list[str], design: Json, page_number: int, numbering: Json) -> None:
    primary = colors.HexColor(str(design["primaryColor"]))
    accent = colors.HexColor(str(design["accentColor"]))
    muted = colors.HexColor(str(design["mutedColor"]))
    style = str(design.get("footerStyle") or "minimal")
    line_one = lines[0] if lines else ""
    line_two = lines[1] if len(lines) > 1 else ""
    line_three = lines[2] if len(lines) > 2 else ""
    page_text = f"Page {page_number}" if numbering.get("enabled", True) else ""

    if style == "band":
        canvas.setFillColor(primary)
        canvas.rect(0, 0, width, 15 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7.2)
        canvas.drawString(left, 9.5 * mm, line_one[:105])
        canvas.drawString(left, 5.5 * mm, line_two[:105])
        canvas.drawRightString(width - right, 7 * mm, (line_three + "  " + page_text).strip()[:105])
    else:
        if style == "double-line":
            canvas.setStrokeColor(primary)
            canvas.setLineWidth(1.2)
            canvas.line(left, 14 * mm, width - right, 14 * mm)
            canvas.setStrokeColor(accent)
            canvas.setLineWidth(0.7)
            canvas.line(left, 12.5 * mm, width - right, 12.5 * mm)
        elif style == "line":
            canvas.setStrokeColor(accent)
            canvas.setLineWidth(1)
            canvas.line(left, 13 * mm, width - right, 13 * mm)
        else:
            canvas.setStrokeColor(colors.HexColor(str(design["borderColor"])))
            canvas.setLineWidth(0.5)
            canvas.line(left, 13 * mm, width - right, 13 * mm)
        canvas.setFillColor(muted)
        canvas.setFont(_pdf_font(str(design["fontFamily"])), 7)
        canvas.drawString(left, 8.5 * mm, (line_one + "  " + line_two).strip()[:130])
        canvas.drawString(left, 5 * mm, line_three[:120])
        if page_text:
            canvas.drawRightString(width - right, 5 * mm, page_text)


def render_pdf(document: Json) -> bytes:
    settings = document.get("settings") or {}
    design = _design(document)
    top, right, bottom, left = _margins(settings)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=_page_size(settings),
        rightMargin=right,
        leftMargin=left,
        topMargin=top + 17 * mm,
        bottomMargin=bottom + 12 * mm,
        title=str((document.get("properties") or {}).get("title") or "Document"),
    )
    styles = getSampleStyleSheet()
    styles["BodyText"].fontName = _pdf_font(str(design["fontFamily"]))
    styles["BodyText"].textColor = colors.HexColor(str(design["textColor"]))
    story = _pdf_story((document.get("body") or {}).get("content") or [], styles, design)
    header_lines = _region_lines(document.get("header"))
    footer_lines = _region_lines(document.get("footer"))
    watermark = (settings.get("watermark") or {}).get("text")
    numbering = settings.get("pageNumbering") or {}

    def decorate(pdf_canvas, _doc):
        pdf_canvas.saveState()
        width, height = _page_size(settings)
        if watermark:
            pdf_canvas.setFillColor(
                colors.Color(
                    0.5,
                    0.5,
                    0.5,
                    alpha=float((settings.get("watermark") or {}).get("opacity", 0.12)),
                )
            )
            pdf_canvas.setFont(_pdf_font(str(design["fontFamily"]), True), 36)
            pdf_canvas.translate(width / 2, height / 2)
            pdf_canvas.rotate(float((settings.get("watermark") or {}).get("rotation", -35)))
            pdf_canvas.drawCentredString(0, 0, str(watermark))
            pdf_canvas.restoreState()
            pdf_canvas.saveState()
        _draw_pdf_header(pdf_canvas, width, height, left, right, header_lines, design)
        _draw_pdf_footer(
            pdf_canvas,
            width,
            left,
            right,
            footer_lines,
            design,
            pdf_canvas.getPageNumber(),
            numbering,
        )
        pdf_canvas.restoreState()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


def _add_page_number(paragraph):
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def _rgb(hex_color: str) -> RGBColor:
    raw = hex_color.lstrip("#")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _docx_add_inline(paragraph, node: Json):
    if node.get("type") == "text":
        run = paragraph.add_run(str(node.get("text", "")))
        for mark in node.get("marks") or []:
            kind = mark.get("type")
            run.bold = run.bold or kind == "bold"
            run.italic = run.italic or kind == "italic"
            run.underline = run.underline or kind == "underline"
        return
    if node.get("type") == "hardBreak":
        paragraph.add_run().add_break()
        return
    for child in node.get("content") or []:
        _docx_add_inline(paragraph, child)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill.lstrip("#"))


def _docx_nodes(container, nodes: list[Json], design: Json):
    primary = _rgb(str(design["primaryColor"]))
    secondary_hex = str(design["secondaryColor"])
    text = _rgb(str(design["textColor"]))
    for node in nodes:
        kind = node.get("type")
        attrs = _attrs(node)
        if kind in {"paragraph", "heading"}:
            style_name = f"Heading {min(3, int(attrs.get('level', 1)))}" if kind == "heading" else None
            paragraph = container.add_paragraph(style=style_name)
            _docx_add_inline(paragraph, node)
            paragraph.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }.get(str(attrs.get("textAlign") or "left"), WD_ALIGN_PARAGRAPH.LEFT)
            if attrs.get("indent"):
                paragraph.paragraph_format.left_indent = Mm(float(attrs["indent"]) * 6)
            for run in paragraph.runs:
                run.font.color.rgb = primary if kind == "heading" else text
                if kind == "heading":
                    run.bold = True
        elif kind == "pageBreak":
            container.add_page_break()
        elif kind == "table":
            rows = node.get("content") or []
            cols = max([len(row.get("content") or []) for row in rows] or [1])
            table = container.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index, cell_node in enumerate(row.get("content") or []):
                    cell = table.cell(row_index, col_index)
                    cell.text = _plain_text(cell_node)
                    if row_index == 0:
                        _shade_cell(cell, str(design["primaryColor"]))
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                run.bold = True
                    elif row_index % 2 == 0:
                        _shade_cell(cell, secondary_hex)
        elif kind == "image":
            raw = _decode_data_uri(str(attrs.get("src") or ""))
            if raw:
                container.add_picture(io.BytesIO(raw), width=Mm(120))
        elif kind == "qrCode":
            qr = qrcode.make(str(attrs.get("value") or attrs.get("label") or "Ithute Document"))
            qbuf = io.BytesIO()
            qr.save(qbuf, format="PNG")
            qbuf.seek(0)
            container.add_picture(qbuf, width=Mm(28))
        elif kind == "textBox":
            table = container.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            _shade_cell(table.cell(0, 0), secondary_hex)
            _docx_nodes(table.cell(0, 0), node.get("content") or [], design)
        elif kind in {"signature", "stamp", "logo", "letterhead", "chart"}:
            paragraph = container.add_paragraph()
            run = paragraph.add_run(str(attrs.get("label") or kind.title()))
            run.bold = True
            run.font.color.rgb = primary
        elif kind in {"bulletList", "orderedList"}:
            for item in node.get("content") or []:
                paragraph = container.add_paragraph(
                    style="List Bullet" if kind == "bulletList" else "List Number"
                )
                _docx_add_inline(paragraph, item)
        elif node.get("content"):
            _docx_nodes(container, node["content"], design)


def _docx_region(container, lines: list[str], design: Json, *, footer: bool = False, numbering: Json | None = None) -> None:
    paragraph = container.paragraphs[0]
    paragraph.clear()
    primary = _rgb(str(design["primaryColor"]))
    muted = _rgb(str(design["mutedColor"]))
    for index, line in enumerate(lines[:3]):
        target = paragraph if index == 0 else container.add_paragraph()
        run = target.add_run(line)
        run.font.name = str(design["fontFamily"])
        run.font.size = Pt(9 if index == 0 and not footer else 7.5)
        run.bold = index == 0 and not footer
        run.font.color.rgb = primary if not footer else muted
    if footer and (numbering or {}).get("enabled", True):
        target = container.add_paragraph()
        target.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = target.add_run("Page ")
        run.font.size = Pt(7.5)
        run.font.color.rgb = muted
        _add_page_number(target)


def render_docx(document: Json) -> bytes:
    out = DocxDocument()
    section = out.sections[0]
    settings = document.get("settings") or {}
    design = _design(document)
    size = _page_size(settings)
    width, height = size
    section.page_width = Mm(width / mm)
    section.page_height = Mm(height / mm)
    if settings.get("orientation") == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
    top, right, bottom, left = [value / mm for value in _margins(settings)]
    section.top_margin = Mm(top)
    section.right_margin = Mm(right)
    section.bottom_margin = Mm(bottom)
    section.left_margin = Mm(left)
    section.header_distance = Mm(float(settings.get("headerDistanceMm") or 8))
    section.footer_distance = Mm(float(settings.get("footerDistanceMm") or 8))

    _docx_region(section.header, _region_lines(document.get("header")), design)
    _docx_region(
        section.footer,
        _region_lines(document.get("footer")),
        design,
        footer=True,
        numbering=settings.get("pageNumbering") or {},
    )

    properties = document.get("properties") or {}
    out.core_properties.title = str(properties.get("title") or "Document")
    out.core_properties.author = str(properties.get("author") or "")
    normal = out.styles["Normal"]
    normal.font.name = str(design["fontFamily"])
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(str(design["textColor"]))
    for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
        heading = out.styles[heading_name]
        heading.font.name = str(design["fontFamily"])
        heading.font.color.rgb = _rgb(str(design["primaryColor"]))

    _docx_nodes(out, (document.get("body") or {}).get("content") or [], design)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()
